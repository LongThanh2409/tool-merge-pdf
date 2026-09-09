from __future__ import annotations

import asyncio
import logging
import os
import re
import shutil
import subprocess
import tempfile
import time
import uuid
from contextlib import asynccontextmanager
from pathlib import Path

import fitz
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
from openpyxl import load_workbook
from openpyxl.worksheet.properties import PageSetupProperties

logging.basicConfig(level=os.getenv("LOG_LEVEL", "INFO"), format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("file-to-pdf")

OFFICE_EXTENSIONS = {".xlsx", ".xls", ".docx", ".doc", ".pptx", ".ppt"}
MIME_ALLOWLIST = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/vnd.ms-excel",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/msword",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation", "application/vnd.ms-powerpoint",
    "application/octet-stream", "application/zip", "",
}
MAX_FILE_SIZE = int(os.getenv("MAX_FILE_SIZE_MB", "100")) * 1024 * 1024
MAX_TOTAL_SIZE = int(os.getenv("MAX_TOTAL_SIZE_MB", "500")) * 1024 * 1024
MAX_FILES = int(os.getenv("MAX_FILES", "50"))
CONVERSION_TIMEOUT = int(os.getenv("CONVERSION_TIMEOUT_SECONDS", "120"))
MAX_CONCURRENCY = int(os.getenv("MAX_CONCURRENCY", "2"))
ALLOWED_ORIGINS = [value.strip() for value in os.getenv("ALLOWED_ORIGINS", "http://localhost:3000").split(",") if value.strip()]
semaphore = asyncio.Semaphore(MAX_CONCURRENCY)


def libreoffice_binary() -> str | None:
    configured = os.getenv("LIBREOFFICE_PATH")
    if configured and Path(configured).is_file():
        return configured
    return shutil.which("libreoffice") or shutil.which("soffice")


def safe_name(filename: str | None, suffix: str) -> str:
    stem = Path(filename or "document").stem
    stem = re.sub(r"[^A-Za-z0-9._-]+", "-", stem).strip(".-")[:80] or "document"
    return f"{stem}-{uuid.uuid4().hex[:12]}{suffix}"


def valid_signature(head: bytes, suffix: str) -> bool:
    if suffix in {".xlsx", ".docx", ".pptx"}:
        return head.startswith(b"PK\x03\x04") or head.startswith(b"PK\x05\x06")
    if suffix in {".xls", ".doc", ".ppt"}:
        return head.startswith(bytes.fromhex("D0CF11E0A1B11AE1"))
    return False


def run_libreoffice(binary: str, source: Path, output: Path, profile: Path, target: str) -> Path:
    output.mkdir()
    profile.mkdir()
    command = [
        binary, "--headless", "--nologo", "--nodefault", "--nolockcheck", "--nofirststartwizard",
        f"-env:UserInstallation=file://{profile}", "--convert-to", target, "--outdir", str(output), str(source),
    ]
    completed = subprocess.run(command, capture_output=True, timeout=CONVERSION_TIMEOUT, check=False)
    if completed.returncode != 0:
        log.warning("CONVERSION_FAILED code=%s", completed.returncode)
        raise RuntimeError("LibreOffice không thể chuyển đổi tài liệu này.")
    suffix = ".pdf" if target.startswith("pdf") else f".{target.split(':', 1)[0]}"
    candidates = list(output.glob(f"*{suffix}"))
    if len(candidates) != 1:
        raise RuntimeError("Không tìm thấy tệp sau khi chuyển đổi.")
    return candidates[0]


def prepare_spreadsheet(source: Path, workdir: Path, binary: str, layout: str, page_size: str, orientation: str, scale_percent: int = 85) -> Path:
    workdir.mkdir(parents=True, exist_ok=True)
    spreadsheet = source
    if source.suffix.lower() == ".xls":
        spreadsheet = run_libreoffice(binary, source, workdir / "xlsx", workdir / "profile-xlsx", "xlsx")
    try:
        workbook = load_workbook(spreadsheet)
        paper_sizes = {"a4": "9", "a3": "8", "letter": "1"}
        for sheet in workbook.worksheets:
            if sheet.max_row == 1 and sheet.max_column == 1 and sheet["A1"].value is None:
                continue
            sheet.print_area = sheet.calculate_dimension()
            if layout == "custom_scale":
                sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=False, autoPageBreaks=True)
                sheet.page_setup.fitToWidth = None
                sheet.page_setup.fitToHeight = None
                sheet.page_setup.scale = scale_percent
            else:
                sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True, autoPageBreaks=False)
                sheet.page_setup.fitToWidth = 1
                sheet.page_setup.fitToHeight = 1 if layout == "single_page" else 0
                sheet.page_setup.scale = None
            sheet.page_setup.paperSize = paper_sizes[page_size]
            resolved_orientation = orientation
            if orientation == "auto":
                total_width = sum(float(sheet.column_dimensions[column].width or 13) for column in sheet.column_dimensions)
                resolved_orientation = "landscape" if sheet.max_column > 8 or total_width > 90 else "portrait"
            sheet.page_setup.orientation = resolved_orientation
        adjusted = workdir / "spreadsheet-adjusted.xlsx"
        workbook.save(adjusted)
        workbook.close()
        return adjusted
    except Exception as exc:
        raise RuntimeError("Không thể điều chỉnh bố cục in của bảng tính.") from exc


def convert_sync(source: Path, workdir: Path, excel_layout: str = "fit_width", page_size: str = "a4", orientation: str = "auto", scale_percent: int = 85) -> tuple[bytes, int]:
    binary = libreoffice_binary()
    if not binary:
        raise RuntimeError("Bộ chuyển đổi LibreOffice chưa được cài đặt trên máy chủ.")
    conversion_source = source
    if source.suffix.lower() in {".xlsx", ".xls"} and excel_layout != "original":
        conversion_source = prepare_spreadsheet(source, workdir, binary, excel_layout, page_size, orientation, scale_percent)
    target = "pdf:calc_pdf_Export" if source.suffix.lower() in {".xlsx", ".xls"} else "pdf"
    pdf_path = run_libreoffice(binary, conversion_source, workdir / "output", workdir / "profile-pdf", target)
    data = pdf_path.read_bytes()
    try:
        document = fitz.open(stream=data, filetype="pdf")
        pages = document.page_count
        document.close()
    except Exception as exc:
        raise RuntimeError("PDF đầu ra không hợp lệ.") from exc
    if not data.startswith(b"%PDF-") or pages < 1:
        raise RuntimeError("PDF đầu ra không hợp lệ.")
    return data, pages


@asynccontextmanager
async def lifespan(_: FastAPI):
    log.info("SERVICE_START libreoffice=%s", bool(libreoffice_binary()))
    yield


app = FastAPI(title="File to PDF Converter", version="1.0.0", lifespan=lifespan)
app.add_middleware(CORSMiddleware, allow_origins=ALLOWED_ORIGINS, allow_credentials=False, allow_methods=["GET", "POST"], allow_headers=["*"])


@app.get("/health")
async def health() -> dict[str, object]:
    ready = libreoffice_binary() is not None
    return {"status": "online" if ready else "error", "converter": "libreoffice", "ready": ready, "timestamp": int(time.time())}


@app.post("/convert")
async def convert(
    file: UploadFile = File(...),
    excel_layout: str = Form("fit_width"),
    page_size: str = Form("a4"),
    orientation: str = Form("auto"),
    scale_percent: int = Form(85),
) -> Response:
    request_id = uuid.uuid4().hex[:12]
    suffix = Path(file.filename or "").suffix.lower()
    log.info("REQUEST id=%s extension=%s", request_id, suffix)
    if suffix not in OFFICE_EXTENSIONS:
        raise HTTPException(415, "Định dạng tệp không được hỗ trợ.")
    if excel_layout not in {"original", "fit_width", "single_page", "custom_scale"} or page_size not in {"a4", "a3", "letter"} or orientation not in {"auto", "portrait", "landscape"} or not 25 <= scale_percent <= 150:
        raise HTTPException(422, "Tùy chọn bố cục không hợp lệ.")
    if (file.content_type or "") not in MIME_ALLOWLIST:
        raise HTTPException(415, "Loại nội dung của tệp không hợp lệ.")
    data = await file.read(MAX_FILE_SIZE + 1)
    await file.close()
    if len(data) > MAX_FILE_SIZE:
        raise HTTPException(413, f"Tệp vượt quá giới hạn {MAX_FILE_SIZE // 1024 // 1024} MB.")
    if not data or not valid_signature(data[:16], suffix):
        raise HTTPException(422, "Tệp có thể đã hỏng hoặc không đúng định dạng.")
    log.info("FILE_RECEIVED id=%s bytes=%s", request_id, len(data))
    async with semaphore:
        try:
            with tempfile.TemporaryDirectory(prefix="file-to-pdf-") as directory:
                workdir = Path(directory)
                source = workdir / safe_name(file.filename, suffix)
                source.write_bytes(data)
                log.info("CONVERSION_START id=%s", request_id)
                pdf, pages = await asyncio.wait_for(
                    asyncio.to_thread(convert_sync, source, workdir, excel_layout, page_size, orientation, scale_percent),
                    timeout=CONVERSION_TIMEOUT + 5,
                )
                log.info("CONVERSION_SUCCESS id=%s pages=%s", request_id, pages)
            log.info("CLEANUP id=%s", request_id)
            return Response(pdf, media_type="application/pdf", headers={"X-PDF-Pages": str(pages), "X-Request-ID": request_id})
        except asyncio.TimeoutError as exc:
            log.error("CONVERSION_FAILED id=%s reason=timeout", request_id)
            raise HTTPException(504, "Chuyển đổi quá thời gian cho phép. Vui lòng thử lại.") from exc
        except RuntimeError as exc:
            log.error("CONVERSION_FAILED id=%s reason=%s", request_id, str(exc))
            raise HTTPException(422, str(exc)) from exc
        except Exception as exc:
            log.exception("ERROR id=%s", request_id)
            raise HTTPException(500, "Không thể chuyển đổi tệp. Vui lòng thử lại.") from exc
