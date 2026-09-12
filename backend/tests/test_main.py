import io
import shutil
import sys
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main
from openpyxl import Workbook, load_workbook

client = TestClient(main.app)


def fake_office_zip() -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
    return output.getvalue()


def test_health_has_explicit_state():
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json()["status"] in {"online", "error"}


def test_health_allows_browser_cors():
    response = client.get("/health", headers={"Origin": "http://localhost:3000"})
    assert response.status_code == 200
    assert response.headers["access-control-allow-origin"] == "*"


def test_rejects_unsupported_extension():
    response = client.post("/convert", files={"file": ("notes.txt", b"hello", "text/plain")})
    assert response.status_code == 415


def test_rejects_bad_signature():
    response = client.post("/convert", files={"file": ("broken.xlsx", b"not-a-zip", "application/octet-stream")})
    assert response.status_code == 422


def test_converts_valid_office_file(monkeypatch):
    sample_pdf = b"%PDF-1.4\n% test"
    monkeypatch.setattr(main, "convert_sync", lambda source, workdir, excel_layout, page_size, orientation, scale_percent: (sample_pdf, 2))
    response = client.post("/convert", files={"file": ("report.xlsx", fake_office_zip(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code == 200
    assert response.content == sample_pdf
    assert response.headers["x-pdf-pages"] == "2"


def test_fit_width_keeps_columns_on_same_page(tmp_path):
    source = tmp_path / "wide.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    for column in range(1, 16):
        sheet.cell(1, column, f"Cột {column}")
        sheet.cell(2, column, column)
    workbook.save(source)
    adjusted = main.prepare_spreadsheet(source, tmp_path / "work", "unused", "fit_width", "a3", "auto")
    result = load_workbook(adjusted)
    output_sheet = result.active
    assert output_sheet.page_setup.fitToWidth == 1
    assert output_sheet.page_setup.fitToHeight == 0
    assert output_sheet.page_setup.orientation == "landscape"
    assert output_sheet.page_setup.paperSize == 8
    assert str(output_sheet.print_area) == "'Sheet'!$A$1:$O$2"
    result.close()


def test_custom_scale_is_written_to_workbook(tmp_path):
    source = tmp_path / "scaled.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "Nội dung"
    workbook.save(source)
    adjusted = main.prepare_spreadsheet(source, tmp_path / "scaled-work", "unused", "custom_scale", "a4", "portrait", 65)
    result = load_workbook(adjusted)
    assert result.active.page_setup.scale == 65
    assert result.active.sheet_properties.pageSetUpPr.fitToPage is False
    result.close()


@pytest.mark.skipif(not (shutil.which("libreoffice") or shutil.which("soffice")), reason="LibreOffice chưa được cài")
def test_libreoffice_is_available_for_integration():
    assert main.libreoffice_binary()
